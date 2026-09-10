# -*- coding: utf-8 -*-
"""
views_solicitudes.py
Modulo de Solicitudes: Login, Creacion de Codigos y Compras.
"""
import json
import re
import os
import io
import uuid
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required, user_passes_test
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_POST
from django.core.mail import send_mail
from django.conf import settings
from django.db import transaction
from django.db.models import Q
from decimal import Decimal
from google import genai
from google.genai import types as genai_types

from taxonomy.models import SKUItem, CodeCreationRequest, PurchaseRequest


def is_admin(user):
    return user.is_active and user.is_staff


def login_view(request):
    error = None
    if request.user.is_authenticated:
        return redirect('taxonomy:solicitudes_dashboard') if is_admin(request.user) else redirect('taxonomy:solicitudes_nueva_codigo')
    if request.method == 'POST':
        user = authenticate(request, username=request.POST.get('username',''), password=request.POST.get('password',''))
        if user:
            login(request, user)
            return redirect('taxonomy:solicitudes_dashboard') if is_admin(user) else redirect('taxonomy:solicitudes_nueva_codigo')
        else:
            error = "Usuario o contrasena incorrectos."
    return render(request, 'solicitudes/login.html', {'error': error})


def logout_view(request):
    logout(request)
    return redirect('taxonomy:login')


@login_required
@user_passes_test(is_admin, login_url='/login/')
def solicitudes_dashboard_view(request):
    status_filter = request.GET.get('status', 'pendiente')
    code_requests = CodeCreationRequest.objects.all()
    purchase_requests = PurchaseRequest.objects.all()
    if status_filter != 'all':
        code_requests = code_requests.filter(status=status_filter)
        purchase_requests = purchase_requests.filter(status=status_filter)

    pending_codes = CodeCreationRequest.objects.filter(status='pendiente').count()
    pending_purchases = PurchaseRequest.objects.filter(status='pendiente').count()
    approved_codes = CodeCreationRequest.objects.filter(status='aprobado').count()
    approved_purchases = PurchaseRequest.objects.filter(status='aprobado').count()
    rejected_codes = CodeCreationRequest.objects.filter(status='rechazado').count()
    rejected_purchases = PurchaseRequest.objects.filter(status='rechazado').count()
    total_codes = CodeCreationRequest.objects.count()
    total_purchases = PurchaseRequest.objects.count()

    context = {
        'code_requests': code_requests,
        'purchase_requests': purchase_requests,
        'status_filter': status_filter,
        'pending_codes_count': pending_codes,
        'pending_purchases_count': pending_purchases,
        'pending_count': pending_codes + pending_purchases,
        'approved_count': approved_codes + approved_purchases,
        'rejected_count': rejected_codes + rejected_purchases,
        'total_codes_count': total_codes,
        'total_purchases_count': total_purchases,
        'total_requests_count': total_codes + total_purchases,
    }
    return render(request, 'solicitudes/dashboard_admin.html', context)


@login_required
@user_passes_test(is_admin, login_url='/login/')
def api_code_request_detail(request, pk):
    """Retorna todos los detalles de una solicitud de código para el modal del Admin."""
    obj = get_object_or_404(CodeCreationRequest, pk=pk)
    base_info = None
    if obj.base_sku:
        base_info = {
            'id': obj.base_sku.id,
            'item_code': obj.base_sku.item_code,
            'item_name': obj.base_sku.item_name,
            'clase': obj.base_sku.clase or '',
            'familia': obj.base_sku.familia or '',
            'subfamilia': obj.base_sku.subfamilia or '',
            'categoria': obj.base_sku.categoria or '',
        }

    file_url = obj.ficha_tecnica_archivo.url if obj.ficha_tecnica_archivo else None
    file_name = obj.ficha_tecnica_archivo.name.rsplit('/', 1)[-1] if obj.ficha_tecnica_archivo else None

    data = {
        'id': obj.pk,
        'solicitante_nombre': obj.solicitante_nombre,
        'grupo_material': obj.grupo_material or '',
        'generated_code': obj.generated_code or '',
        'proposed_description': obj.proposed_description or '',
        'justification': obj.justification or '',
        'proveedor_nombre': obj.proveedor_nombre or '',
        'codigo_catalogo_proveedor': obj.codigo_catalogo_proveedor or '',
        'es_importado': obj.es_importado,
        'moneda_precio': obj.moneda_precio or 'CLP',
        'precio_referencial': str(obj.precio_referencial) if obj.precio_referencial is not None else '',
        'lote_minimo': obj.lote_minimo or '',
        'unidad_empaque': obj.unidad_empaque or '',
        'incoterm': obj.incoterm or 'NA',
        'ficha_tecnica': obj.ficha_tecnica or '',
        'ficha_tecnica_archivo_url': file_url,
        'ficha_tecnica_archivo_name': file_name,
        'clase_propuesta': obj.clase_propuesta or '',
        'familia_propuesta': obj.familia_propuesta or '',
        'subfamilia_propuesta': obj.subfamilia_propuesta or '',
        'categoria_propuesta': obj.categoria_propuesta or '',
        'base_sku': base_info,
        'status': obj.status,
        'status_display': obj.get_status_display(),
        'admin_notes': obj.admin_notes or '',
        'created_at': obj.created_at.strftime('%d/%m/%Y %H:%M'),
    }
    return JsonResponse({'success': True, 'data': data})


@login_required
@user_passes_test(is_admin, login_url='/login/')
def approve_code_request(request, pk):
    """
    Permite al Admin enriquecer datos de la solicitud y decidir si:
    - 'approve': Aprueba y crea automáticamente el artículo en el Maestro (SKUItem).
    - 'save': Guarda cambios de edición sin aprobar aún (mantiene estado).
    - 'reject': Rechaza la solicitud con justificación.
    """
    if request.method != 'POST':
        return JsonResponse({'success': False}, status=405)
    obj = get_object_or_404(CodeCreationRequest, pk=pk)
    data = request.POST

    action = data.get('action', 'save')

    # Actualización de datos enriquecidos por el administrador
    if 'generated_code' in data:
        obj.generated_code = data.get('generated_code', '').strip() or None
    if 'proposed_description' in data:
        obj.proposed_description = data.get('proposed_description', '').strip()
    if 'grupo_material' in data:
        obj.grupo_material = data.get('grupo_material', '').strip() or None
    if 'clase_propuesta' in data:
        obj.clase_propuesta = data.get('clase_propuesta', '').strip() or None
    if 'familia_propuesta' in data:
        obj.familia_propuesta = data.get('familia_propuesta', '').strip() or None
    if 'subfamilia_propuesta' in data:
        obj.subfamilia_propuesta = data.get('subfamilia_propuesta', '').strip() or None
    if 'categoria_propuesta' in data:
        obj.categoria_propuesta = data.get('categoria_propuesta', '').strip() or None
    if 'proveedor_nombre' in data:
        obj.proveedor_nombre = data.get('proveedor_nombre', '').strip() or None
    if 'codigo_catalogo_proveedor' in data:
        obj.codigo_catalogo_proveedor = data.get('codigo_catalogo_proveedor', '').strip() or None
    if 'unidad_empaque' in data:
        obj.unidad_empaque = data.get('unidad_empaque', '').strip() or None
    if 'ficha_tecnica' in data:
        obj.ficha_tecnica = data.get('ficha_tecnica', '').strip() or None

    lote_raw = data.get('lote_minimo', '').strip() if 'lote_minimo' in data else None
    if lote_raw is not None:
        try:
            obj.lote_minimo = int(lote_raw) if lote_raw else None
        except ValueError:
            pass

    precio_raw = data.get('precio_referencial', '').strip() if 'precio_referencial' in data else None
    if precio_raw is not None:
        try:
            obj.precio_referencial = Decimal(precio_raw) if precio_raw else None
        except Exception:
            pass

    if 'admin_notes' in data:
        obj.admin_notes = data.get('admin_notes', '').strip()

    sku_created = False
    new_sku_code = None

    if action == 'approve':
        obj.status = 'aprobado'
        # ── CREACIÓN AUTOMÁTICA EN EL MAESTRO DE MATERIALES (SKUItem) ──
        code = (obj.generated_code or '').strip()
        if code:
            sku, _ = SKUItem.objects.get_or_create(item_code=code)
            sku.item_name = obj.proposed_description or ''
            sku.nombre_grupo = obj.grupo_material or ''
            sku.clase = obj.clase_propuesta or ''
            sku.familia = obj.familia_propuesta or ''
            sku.subfamilia = obj.subfamilia_propuesta or ''
            sku.categoria = obj.categoria_propuesta or ''
            if obj.precio_referencial:
                sku.costo_un = obj.precio_referencial
                sku.moneda = obj.moneda_precio or 'CLP'
            sku.check_incomplete()
            sku.save()
            sku_created = True
            new_sku_code = sku.item_code
        _enviar_correo_aprobacion(obj)
    elif action == 'reject':
        obj.status = 'rechazado'
        _enviar_correo_rechazo(obj)

    obj.save()

    return JsonResponse({
        'success': True,
        'new_status': obj.get_status_display(),
        'status': obj.status,
        'sku_created': sku_created,
        'new_sku_code': new_sku_code
    })


def _enviar_correo_aprobacion(solicitud):
    dest = getattr(settings, 'ABASTECIMIENTO_EMAIL', 'abastecimiento@pesco.cl')
    try:
        send_mail(
            subject=f"[PESCO] Solicitud de Código #{solicitud.pk} APROBADA - Código SAP: {solicitud.generated_code}",
            message=(
                f"La solicitud de creación de código SAP #{solicitud.pk} ha sido APROBADA y dada de alta en el Maestro.\n\n"
                f"Código SAP Asignado: {solicitud.generated_code}\n"
                f"Descripción Definitiva: {solicitud.proposed_description}\n"
                f"Solicitante: {solicitud.solicitante_nombre}\n"
                f"Grupo de Material: {solicitud.grupo_material or '-'}\n"
                f"Clase: {solicitud.clase_propuesta or '-'}\n"
                f"Familia: {solicitud.familia_propuesta or '-'}\n"
                f"Subfamilia: {solicitud.subfamilia_propuesta or '-'}\n"
                f"Proveedor: {solicitud.proveedor_nombre or '-'}\n"
                f"Cod. Catálogo Proveedor: {solicitud.codigo_catalogo_proveedor or '-'}\n\n"
                f"Observaciones Admin: {solicitud.admin_notes or 'Sin observaciones.'}\n\n"
                f"Portal de Gestión: http://127.0.0.1:8000/solicitudes/"
            ),
            from_email=settings.DEFAULT_FROM_EMAIL,
            recipient_list=[dest],
            fail_silently=True,
        )
    except Exception:
        pass


def _enviar_correo_rechazo(solicitud):
    dest = getattr(settings, 'ABASTECIMIENTO_EMAIL', 'abastecimiento@pesco.cl')
    try:
        send_mail(
            subject=f"[PESCO] Solicitud de Código #{solicitud.pk} RECHAZADA",
            message=(
                f"La solicitud de creación de código SAP #{solicitud.pk} ha sido RECHAZADA.\n\n"
                f"Descripción solicitada: {solicitud.proposed_description}\n"
                f"Solicitante: {solicitud.solicitante_nombre}\n\n"
                f"Motivo / Observaciones del Administrador:\n{solicitud.admin_notes or 'No se especificó motivo.'}\n\n"
                f"Portal de Gestión: http://127.0.0.1:8000/solicitudes/"
            ),
            from_email=settings.DEFAULT_FROM_EMAIL,
            recipient_list=[dest],
            fail_silently=True,
        )
    except Exception:
        pass


@login_required
@user_passes_test(is_admin, login_url='/login/')
def approve_purchase_request(request, pk):
    if request.method != 'POST':
        return JsonResponse({'success': False}, status=405)
    obj = get_object_or_404(PurchaseRequest, pk=pk)
    action = request.POST.get('action')
    obj.admin_notes = request.POST.get('admin_notes', '').strip()
    if action == 'approve':
        obj.status = 'aprobado'
    elif action == 'reject':
        obj.status = 'rechazado'
    obj.save()
    return JsonResponse({'success': True, 'new_status': obj.get_status_display()})



@login_required
def nueva_solicitud_codigo_view(request):
    # Enviar al template los grupos de material disponibles en el Maestro
    if request.method == 'GET':
        grupos = list(
            SKUItem.objects.exclude(nombre_grupo__isnull=True)
            .exclude(nombre_grupo='')
            .values_list('nombre_grupo', flat=True)
            .distinct()
            .order_by('nombre_grupo')
        )
        return render(request, 'solicitudes/form_creacion_codigo.html', {'grupos_material': grupos})

    # ── POST: guardar solicitud ────────────────────────────────────────────────
    data = request.POST
    files = request.FILES

    base_sku = None
    base_sku_id = data.get('base_sku_id')
    if base_sku_id:
        try:
            base_sku = SKUItem.objects.get(pk=int(base_sku_id))
        except (SKUItem.DoesNotExist, ValueError):
            pass

    # Precio referencial (puede venir vacío)
    precio_ref = None
    precio_raw = data.get('precio_referencial', '').strip()
    if precio_raw:
        try:
            precio_ref = Decimal(precio_raw.replace(',', '.'))
        except Exception:
            pass

    # Lote mínimo (puede venir vacío)
    lote_min = None
    lote_raw = data.get('lote_minimo', '').strip()
    if lote_raw:
        try:
            lote_min = int(lote_raw)
        except Exception:
            pass

    solicitud = CodeCreationRequest(
        base_sku=base_sku,
        generated_code=data.get('generated_code', '').strip() or None,
        solicitante_nombre=data.get('solicitante_nombre', '').strip(),
        grupo_material=data.get('grupo_material', '').strip() or None,
        proposed_description=data.get('proposed_description', '').strip(),
        justification=data.get('justification', '').strip(),
        # Proveedor
        proveedor_nombre=data.get('proveedor_nombre', '').strip() or None,
        codigo_catalogo_proveedor=data.get('codigo_catalogo_proveedor', '').strip() or None,
        es_importado=data.get('es_importado') == 'on',
        moneda_precio=data.get('moneda_precio', 'CLP') or 'CLP',
        precio_referencial=precio_ref,
        lote_minimo=lote_min,
        unidad_empaque=data.get('unidad_empaque', '').strip() or None,
        incoterm=data.get('incoterm', 'NA') or 'NA',
        # Ficha técnica
        ficha_tecnica=data.get('ficha_tecnica', '').strip() or None,
        # Taxonomía
        clase_propuesta=data.get('clase_propuesta', '').strip() or None,
        familia_propuesta=data.get('familia_propuesta', '').strip() or None,
        subfamilia_propuesta=data.get('subfamilia_propuesta', '').strip() or None,
        categoria_propuesta=data.get('categoria_propuesta', '').strip() or None,
    )

    # Archivo adjunto (PDF / JPG / PNG)
    archivo = files.get('ficha_tecnica_archivo')
    if archivo:
        ext = archivo.name.rsplit('.', 1)[-1].lower()
        if ext not in ('pdf', 'jpg', 'jpeg', 'png'):
            return JsonResponse({'success': False, 'error': 'Formato de archivo no válido. Use PDF, JPG o PNG.'}, status=400)
        solicitud.ficha_tecnica_archivo = archivo

    solicitud.save()
    _enviar_correo_creacion(solicitud)
    return JsonResponse({'success': True, 'id': solicitud.pk})


def _enviar_correo_creacion(solicitud):
    dest = getattr(settings, 'ABASTECIMIENTO_EMAIL', 'abastecimiento@pesco.cl')
    base_info = (
        f"Codigo base referencia: {solicitud.base_sku.item_code} - {solicitud.base_sku.item_name}"
        if solicitud.base_sku else "Sin codigo base"
    )
    importado_str = "SÍ" if solicitud.es_importado else "No"
    proveedor_block = (
        f"\n── PROVEEDOR ──\n"
        f"Proveedor: {solicitud.proveedor_nombre or '-'}\n"
        f"Cod. Catálogo Proveedor: {solicitud.codigo_catalogo_proveedor or '-'}\n"
        f"Artículo importado: {importado_str}\n"
        f"Precio referencial: {solicitud.moneda_precio} {solicitud.precio_referencial or '-'}\n"
        f"Lote mínimo: {solicitud.lote_minimo or '-'} uds\n"
        f"Unidad de empaque: {solicitud.unidad_empaque or '-'}\n"
        f"Incoterm: {solicitud.incoterm}\n"
    )
    ficha_block = (
        f"\n── FICHA TÉCNICA ──\n{solicitud.ficha_tecnica}"
        if solicitud.ficha_tecnica else ""
    )
    archivo_block = (
        f"\nArchivo adjunto: {solicitud.ficha_tecnica_archivo.name}"
        if solicitud.ficha_tecnica_archivo else ""
    )
    try:
        send_mail(
            subject=f"[PESCO] Nueva Solicitud de Código #{solicitud.pk} - {solicitud.solicitante_nombre}",
            message=(
                f"Nueva solicitud de creación de código SAP.\n\n"
                f"ID: #{solicitud.pk}\n"
                f"Solicitante: {solicitud.solicitante_nombre}\n"
                f"Grupo de Material: {solicitud.grupo_material or '-'}\n"
                f"Código propuesto: {solicitud.generated_code or 'Pendiente'}\n"
                f"Descripción: {solicitud.proposed_description}\n"
                f"Motivo: {solicitud.justification}\n"
                f"{base_info}"
                f"{proveedor_block}"
                f"{ficha_block}"
                f"{archivo_block}\n\n"
                f"Portal: http://127.0.0.1:8000/solicitudes/"
            ),
            from_email=settings.DEFAULT_FROM_EMAIL,
            recipient_list=[dest],
            fail_silently=True,
        )
    except Exception:
        pass


@login_required
def nueva_solicitud_compra_view(request):
    if request.method == 'POST':
        data = request.POST
        try:
            cant = Decimal(str(data.get('cantidad_solicitada','0') or '0'))
            cu = Decimal(str(data.get('costo_unitario','0') or '0'))
            ct = cant * cu
        except Exception:
            cant, cu, ct = 0, 0, 0
        solicitud = PurchaseRequest.objects.create(
            solicitante_nombre=data.get('solicitante_nombre','').strip(),
            proveedor=data.get('proveedor','').strip() or None,
            codigo_compra=data.get('codigo_compra','').strip() or None,
            descripcion=data.get('descripcion','').strip(),
            cantidad_solicitada=cant,
            costo_unitario=cu,
            costo_total=ct,
            tipo_compra=data.get('tipo_compra','stock'),
            justification=data.get('justification','').strip() or None,
        )
        dest = getattr(settings, 'ABASTECIMIENTO_EMAIL', 'abastecimiento@pesco.cl')
        try:
            send_mail(
                subject=f"[PESCO] Nueva Solicitud de Compra #{solicitud.pk} - {solicitud.solicitante_nombre}",
                message=(
                    f"Nueva solicitud de compra.\n\n"
                    f"ID: #{solicitud.pk}\nSolicitante: {solicitud.solicitante_nombre}\n"
                    f"Proveedor: {solicitud.proveedor or '-'}\nCodigo: {solicitud.codigo_compra or '-'}\n"
                    f"Descripcion: {solicitud.descripcion}\nCantidad: {solicitud.cantidad_solicitada}\n"
                    f"Costo Unit.: ${solicitud.costo_unitario}\nCosto Total: ${solicitud.costo_total}\n"
                    f"Tipo: {solicitud.get_tipo_compra_display()}\n\nPortal: http://127.0.0.1:8000/solicitudes/"
                ),
                from_email=settings.DEFAULT_FROM_EMAIL,
                recipient_list=[dest],
                fail_silently=True,
            )
        except Exception:
            pass
        return JsonResponse({'success': True, 'id': solicitud.pk})
    return render(request, 'solicitudes/form_compra.html')


def api_next_correlative(request, base_sku_code):
    try:
        base_sku = get_object_or_404(SKUItem, item_code=base_sku_code)
        code = base_sku.item_code.strip()
        match = re.match(r'^([A-Za-z\-_\.]+)(\d+)$', code)
        if not match:
            match = re.match(r'^(\d{1,6})(\d{4,})$', code)
        if match:
            prefix = match.group(1)
            num_part = match.group(2)
            num_len = len(num_part)
            similar = SKUItem.objects.filter(item_code__startswith=prefix)
            max_num = 0
            for sku in similar:
                m2 = re.match(rf'^{re.escape(prefix)}(\d+)$', sku.item_code.strip())
                if m2:
                    n = int(m2.group(1))
                    if n > max_num:
                        max_num = n
            next_code = f"{prefix}{str(max_num + 1).zfill(num_len)}"
        else:
            next_code = f"{code}_NUEVO"
        return JsonResponse({
            'success': True,
            'base_code': code,
            'next_code': next_code,
            'base_item_name': base_sku.item_name,
            'base_clase': base_sku.clase or '',
            'base_familia': base_sku.familia or '',
            'base_subfamilia': base_sku.subfamilia or '',
            'base_categoria': base_sku.categoria or '',
        })
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


def api_buscar_skus(request):
    q = request.GET.get('q', '').strip()
    if len(q) < 2:
        return JsonResponse({'results': []})
    qs = SKUItem.objects.filter(
        Q(item_code__icontains=q) | Q(item_name__icontains=q)
    ).values('id', 'item_code', 'item_name', 'clase', 'familia', 'subfamilia', 'categoria')[:15]
    return JsonResponse({'results': list(qs)})


@require_POST
def api_ai_assistant(request):
    try:
        body = json.loads(request.body)
        descripcion = body.get('descripcion', '').strip()
        justificacion = body.get('justificacion', '').strip()
        base_item_name = body.get('base_item_name', '').strip()
        grupo_material = body.get('grupo_material', '').strip()
        tipo = body.get('tipo', 'codigo')

        api_key = os.getenv('GOOGLE_API_KEY') or os.getenv('GEMINI_API_KEY')
        if not api_key:
            return JsonResponse({'success': False, 'error': 'API Key no configurada.'}, status=500)

        client = genai.Client(api_key=api_key)

        if tipo == 'codigo':
            # ── Extraer palabras clave para la búsqueda inteligente de contexto
            stop_words = {'CON', 'PARA', 'POR', 'SIN', 'DEL', 'LOS', 'LAS', 'UNA', 'UNO', 'DE', 'EL', 'LA', 'ANTE', 'BAJO', 'CADA'}
            words = [w.strip().upper() for w in descripcion.split() if len(w.strip()) > 2 and w.strip().upper() not in stop_words]

            skus_contexto = []
            seen_codes = set()

            # Prioridad 1: Coincidencia por palabras clave dentro del mismo grupo
            if grupo_material and words:
                q_kw = Q()
                for w in words:
                    q_kw |= Q(item_name__icontains=w)
                p1 = SKUItem.objects.filter(Q(nombre_grupo__iexact=grupo_material) & q_kw).values(
                    'item_code', 'item_name', 'clase', 'familia', 'subfamilia', 'nombre_grupo'
                )[:6]
                for item in p1:
                    if item['item_code'] not in seen_codes:
                        skus_contexto.append(item)
                        seen_codes.add(item['item_code'])

            # Prioridad 2: Coincidencia por palabras clave a nivel todo el Maestro
            if words and len(skus_contexto) < 8:
                q_kw = Q()
                for w in words:
                    q_kw |= Q(item_name__icontains=w)
                p2 = SKUItem.objects.filter(q_kw).exclude(item_code__in=seen_codes).values(
                    'item_code', 'item_name', 'clase', 'familia', 'subfamilia', 'nombre_grupo'
                )[:6]
                for item in p2:
                    if item['item_code'] not in seen_codes:
                        skus_contexto.append(item)
                        seen_codes.add(item['item_code'])

            # Prioridad 3: Artículos generales del mismo grupo
            if grupo_material and len(skus_contexto) < 10:
                p3 = SKUItem.objects.filter(nombre_grupo__iexact=grupo_material).exclude(
                    item_code__in=seen_codes
                ).values('item_code', 'item_name', 'clase', 'familia', 'subfamilia', 'nombre_grupo')[:6]
                for item in p3:
                    if item['item_code'] not in seen_codes:
                        skus_contexto.append(item)
                        seen_codes.add(item['item_code'])

            contexto_str = "\n".join(
                f"  - [{s['item_code']}] {s['item_name']} (Grupo: {s.get('nombre_grupo') or '-'}, Clase: {s.get('clase') or '-'})"
                for s in skus_contexto
            ) if skus_contexto else "  (No se encontraron artículos similares en el maestro)"

            prompt = f"""Eres un experto en catalogos de maestros de articulos para PESCO S.A., empresa industrial.
Tu tarea es: (1) mejorar la descripcion del nuevo articulo al estandar SAP de PESCO, y (2) sugerir el codigo SAP existente mas adecuado como referencia base para derivar el correlativo numérico.

Grupo de material solicitante: "{grupo_material or 'No especificado'}"
Descripcion ingresada: "{descripcion}"
Justificacion: "{justificacion}"
Articulo base seleccionado manualmente: "{base_item_name or 'Ninguno'}"

Codigos existentes en el Maestro para este grupo/descripcion:
{contexto_str}

Instrucciones OBLIGATORIAS:
- "descripcion_mejorada": Estandariza en MAYUSCULAS, concisa (max 100 caracteres), formato SAP.
- "codigo_referencia_sugerido": DEBES ELEGIR el item_code del listado de contexto que sea mas similar funcionalmente, semánticamente o por categoría al articulo solicitado (dando preferencia a los articulos del mismo grupo "{grupo_material}"). SE REQUIERE UN CODIGO DE REFERENCIA PARA GENERAR EL CORRELATIVO SAP. Solo devuelve null si el listado de contexto esta 100% vacio.
- "nombre_referencia_sugerido": El item_name del codigo sugerido. Si null, devuelve null.
- "sugerencias": Lista de 2-3 consejos para completar mejor la solicitud.
- "advertencias": Alertas si falta info critica o hay riesgo de duplicado.
- "completitud_score": Porcentaje 0-100 de que tan completa esta la solicitud.

Responde SOLO en JSON sin texto adicional:
{{
  "descripcion_mejorada": "...",
  "codigo_referencia_sugerido": "20701539",
  "nombre_referencia_sugerido": "NOMBRE DEL ARTICULO REFERENCIA",
  "sugerencias": ["...", "..."],
  "advertencias": ["..."],
  "completitud_score": 85
}}"""
        else:
            prompt = f"""Eres un experto en compras y abastecimiento industrial para PESCO S.A.
Analiza y mejora esta solicitud de compra.

Articulo: "{descripcion}"
Justificacion: "{justificacion}"

Responde SOLO en JSON sin texto adicional:
{{
  "descripcion_mejorada": "descripcion tecnica clara del articulo",
  "sugerencias": ["consejo 1", "consejo 2"],
  "advertencias": ["advertencia si falta info critica"],
  "completitud_score": 85
}}"""

        models_to_try = ['gemini-3.6-flash', 'gemini-2.5-flash', 'gemini-1.5-flash']
        response = None
        last_err = None
        for m in models_to_try:
            try:
                response = client.models.generate_content(
                    model=m,
                    contents=prompt,
                    config=genai_types.GenerateContentConfig(
                        response_mime_type="application/json",
                        temperature=0.25,
                    )
                )
                break
            except Exception as err:
                last_err = err

        if not response:
            raise Exception(f"No se pudo contactar Gemini API: {last_err}")

        result = json.loads(response.text)

        # Fallback determinista si Gemini devolvió null pero hay candidatos de contexto
        if tipo == 'codigo' and not result.get('codigo_referencia_sugerido') and skus_contexto:
            top_sku = skus_contexto[0]
            result['codigo_referencia_sugerido'] = top_sku['item_code']
            result['nombre_referencia_sugerido'] = top_sku['item_name']

        return JsonResponse({'success': True, **result})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


# ══════════════════════════════════════════════════════════════════════════════
# ── BANDEJA DE CARGA MASIVA & IMPORTACIÓN EXCEL (ARQUITECTURA SENIOR) ──────────
# ══════════════════════════════════════════════════════════════════════════════

def api_descargar_plantilla(request, tipo):
    """
    Genera y retorna un archivo Excel (.xlsx) oficial de PESCO como plantilla descargable.
    """
    wb = openpyxl.Workbook()
    ws = wb.active

    # Dark header fill & styles PESCO
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    if tipo == 'codigo':
        ws.title = "Carga Masiva Codigos"
        headers = [
            "DESCRIPCION", "GRUPO_MATERIAL", "PROVEEDOR", "COD_CATALOGO_PROVEEDOR",
            "UNIDAD_EMPAQUE", "LOTE_MINIMO", "PRECIO_REFERENCIAL", "ES_IMPORTADO",
            "MOTIVO_JUSTIFICACION", "ESPECIFICACIONES_TECNICAS"
        ]
        sample_rows = [
            [
                "GRUA ARTICULADA ING 8500C 3H KM", "EQUIPOS IZAJE", "ING Cranes", "ING-8500C-KM",
                "1 UN", 1, 45000, "SI", "Alta masiva equipos mineros proyecto 2026",
                "Grúa articulada capacidad 8.5 ton con kit minero completo."
            ],
            [
                "BASURERO PAPELERO CON PEDAL 5L ACERO INOXIDABLE", "PESCO AMBULANCIAS", "Inversiones J&A", "MKRI57L805-1",
                "1 UN", 1, 15000, "NO", "Equipamiento para ambulancias de rescate",
                "Basurero acero inox 202 con pedal reforzado."
            ]
        ]
        filename = "Plantilla_Carga_Masiva_Codigos_PESCO.xlsx"
    else:
        ws.title = "Carga Masiva Compras"
        headers = [
            "CODIGO_SAP", "DESCRIPCION", "CANTIDAD", "COSTO_UNITARIO",
            "PROVEEDOR", "TIPO_COMPRA", "JUSTIFICACION"
        ]
        sample_rows = [
            [
                "13171071", "GRUA ARTICULADA ING 8500C 3H KM", 2, 45000,
                "ING Cranes", "Stock", "Compra masiva para stock pañol central"
            ],
            [
                "20701539", "PAPELERO AC.INOX 3 LTS C/PEDAL", 10, 12000,
                "Inversiones J&A", "Calzada", "Compra calzada para orden de trabajo OF-2607813"
            ]
        ]
        filename = "Plantilla_Carga_Masiva_Compras_PESCO.xlsx"

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    for r in sample_rows:
        ws.append(r)

    ws.row_dimensions[1].height = 28
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@require_POST
def api_parse_excel_batch(request, tipo):
    """
    Parsea una planilla Excel (.xlsx) subida por el usuario y retorna el listado de ítems
    validados y enriquecidos para la Bandeja Masiva.
    """
    excel_file = request.FILES.get('excel_file')
    if not excel_file:
        return JsonResponse({'success': False, 'error': 'No se adjuntó archivo Excel.'}, status=400)

    try:
        wb = openpyxl.load_workbook(excel_file, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))

        if not rows or len(rows) < 2:
            return JsonResponse({'success': False, 'error': 'El archivo Excel está vacío o no contiene filas de datos.'}, status=400)

        # Header normalization
        raw_headers = [str(cell).strip().upper() if cell is not None else '' for cell in rows[0]]
        header_map = {}
        for idx, h in enumerate(raw_headers):
            if h:
                header_map[h] = idx

        items = []
        for r_idx, row in enumerate(rows[1:], start=2):
            if not any(row):
                continue

            def get_val(key_list, default=''):
                for k in key_list:
                    if k in header_map and header_map[k] < len(row):
                        v = row[header_map[k]]
                        if v is not None:
                            return str(v).strip()
                return default

            if tipo == 'codigo':
                desc = get_val(['DESCRIPCION', 'DESCRIPCION_ARTICULO', 'ITEM_NAME', 'NOMBRE'])
                if not desc:
                    continue

                grupo = get_val(['GRUPO_MATERIAL', 'GRUPO', 'AREA', 'NOMBRE_GRUPO'])
                proveedor = get_val(['PROVEEDOR', 'NOMBRE_PROVEEDOR', 'PROVEEDOR_NOMBRE'])
                cod_cat = get_val(['COD_CATALOGO_PROVEEDOR', 'CODIGO_CATALOGO_PROVEEDOR', 'COD_CATALOGO', 'CATALOGO'])
                empaque = get_val(['UNIDAD_EMPAQUE', 'EMPAQUE', 'UNIDAD'])
                lote_raw = get_val(['LOTE_MINIMO', 'LOTE_MIN', 'LOTE'])
                precio_raw = get_val(['PRECIO_REFERENCIAL', 'PRECIO_REF', 'PRECIO', 'COSTO'])
                importado_raw = get_val(['ES_IMPORTADO', 'IMPORTADO', 'INTERNACIONAL']).upper()
                just = get_val(['MOTIVO_JUSTIFICACION', 'JUSTIFICACION', 'MOTIVO'])
                ficha = get_val(['ESPECIFICACIONES_TECNICAS', 'FICHA_TECNICA', 'ESPECIFICACIONES'])

                es_importado = importado_raw in ('SI', 'SÍ', 'TRUE', '1', 'YES')

                # Trazar correlativo / sugerencia si aplica
                generated_code = ""
                base_sku_code = ""
                if desc:
                    keywords = [w for w in desc.split() if len(w) > 3][:2]
                    q_kw = Q()
                    for kw in keywords:
                        q_kw |= Q(item_name__icontains=kw)
                    candidate = SKUItem.objects.filter(q_kw).first()
                    if candidate:
                        base_sku_code = candidate.item_code
                        m = re.match(r'^([A-Za-z\-_\.]+)(\d+)$', candidate.item_code)
                        if not m:
                            m = re.match(r'^(\d{1,6})(\d{4,})$', candidate.item_code)
                        if m:
                            prefix = m.group(1)
                            num_part = m.group(2)
                            num_len = len(num_part)
                            similar = SKUItem.objects.filter(item_code__startswith=prefix)
                            max_n = 0
                            for s in similar:
                                m2 = re.match(rf'^{re.escape(prefix)}(\d+)$', s.item_code.strip())
                                if m2 and int(m2.group(1)) > max_n:
                                    max_n = int(m2.group(1))
                            generated_code = f"{prefix}{str(max_n + 1).zfill(num_len)}"

                items.append({
                    'tipo': 'codigo',
                    'proposed_description': desc,
                    'grupo_material': grupo,
                    'generated_code': generated_code,
                    'base_sku_code': base_sku_code,
                    'proveedor_nombre': proveedor,
                    'codigo_catalogo_proveedor': cod_cat,
                    'unidad_empaque': empaque,
                    'lote_minimo': lote_raw,
                    'precio_referencial': precio_raw,
                    'es_importado': es_importado,
                    'justification': just,
                    'ficha_tecnica': ficha,
                })
            else: # tipo == 'compra'
                desc = get_val(['DESCRIPCION', 'DESCRIPCION_ARTICULO', 'ITEM_NAME'])
                cod_compra = get_val(['CODIGO_SAP', 'CODIGO_COMPRA', 'CODIGO_ITEM', 'SKU'])
                if not desc and not cod_compra:
                    continue

                cant_raw = get_val(['CANTIDAD', 'CANTIDAD_SOLICITADA', 'CANT'], '1')
                cu_raw = get_val(['COSTO_UNITARIO', 'COSTO_UNIT', 'PRECIO_UNITARIO'], '0')
                prov = get_val(['PROVEEDOR', 'NOMBRE_PROVEEDOR'])
                tipo_c = get_val(['TIPO_COMPRA', 'TIPO'], 'stock').lower()
                just = get_val(['JUSTIFICACION', 'MOTIVO'])

                if tipo_c not in ('stock', 'calzada'):
                    tipo_c = 'stock'

                items.append({
                    'tipo': 'compra',
                    'codigo_compra': cod_compra,
                    'descripcion': desc or f"Compra artículo {cod_compra}",
                    'cantidad_solicitada': cant_raw,
                    'costo_unitario': cu_raw,
                    'proveedor': prov,
                    'tipo_compra': tipo_c,
                    'justification': just,
                })

        return JsonResponse({'success': True, 'items': items, 'total': len(items)})
    except Exception as e:
        return JsonResponse({'success': False, 'error': f"Error al leer el archivo Excel: {str(e)}"}, status=400)


@require_POST
def api_procesar_lote_masivo(request):
    """
    Procesa de forma atómica (transaccional) la lista completa de ítems de la Bandeja Masiva.
    """
    try:
        body = json.loads(request.body)
        solicitante = body.get('solicitante_nombre', '').strip() or 'Usuario Sistema'
        tipo = body.get('tipo', 'codigo')
        items = body.get('items', [])

        if not items:
            return JsonResponse({'success': False, 'error': 'El lote de la bandeja está vacío.'}, status=400)

        lote_id = f"LOTE-{uuid.uuid4().hex[:8].upper()}"
        created_records = []

        with transaction.atomic():
            if tipo == 'codigo':
                for it in items:
                    desc = str(it.get('proposed_description', '')).strip()
                    if not desc:
                        continue
                    precio_ref = None
                    p_raw = str(it.get('precio_referencial', '')).strip()
                    if p_raw:
                        try: precio_ref = Decimal(p_raw)
                        except Exception: pass

                    lote_min = None
                    l_raw = str(it.get('lote_minimo', '')).strip()
                    if l_raw:
                        try: lote_min = int(l_raw)
                        except Exception: pass

                    code_gen = str(it.get('generated_code', '')).strip() or None

                    rec = CodeCreationRequest.objects.create(
                        solicitante_nombre=solicitante,
                        grupo_material=str(it.get('grupo_material', '')).strip() or None,
                        generated_code=code_gen,
                        proposed_description=desc,
                        justification=str(it.get('justification', '')).strip() or "Carga masiva en lote",
                        proveedor_nombre=str(it.get('proveedor_nombre', '')).strip() or None,
                        codigo_catalogo_proveedor=str(it.get('codigo_catalogo_proveedor', '')).strip() or None,
                        unidad_empaque=str(it.get('unidad_empaque', '')).strip() or None,
                        lote_minimo=lote_min,
                        precio_referencial=precio_ref,
                        es_importado=bool(it.get('es_importado')),
                        ficha_tecnica=str(it.get('ficha_tecnica', '')).strip() or None,
                        lote_id=lote_id,
                        status='pendiente'
                    )
                    created_records.append(rec)
            else: # tipo == 'compra'
                for it in items:
                    desc = str(it.get('descripcion', '')).strip()
                    if not desc:
                        continue
                    cant = Decimal('1')
                    cu = Decimal('0')
                    try: cant = Decimal(str(it.get('cantidad_solicitada', '1')))
                    except Exception: pass
                    try: cu = Decimal(str(it.get('costo_unitario', '0')))
                    except Exception: pass
                    ct = cant * cu

                    rec = PurchaseRequest.objects.create(
                        solicitante_nombre=solicitante,
                        proveedor=str(it.get('proveedor', '')).strip() or None,
                        codigo_compra=str(it.get('codigo_compra', '')).strip() or None,
                        descripcion=desc,
                        cantidad_solicitada=cant,
                        costo_unitario=cu,
                        costo_total=ct,
                        tipo_compra=str(it.get('tipo_compra', 'stock')).lower(),
                        justification=str(it.get('justification', '')).strip() or "Carga masiva en lote",
                        lote_id=lote_id,
                        status='pendiente'
                    )
                    created_records.append(rec)

        _enviar_correo_lote_masivo(solicitante, tipo, lote_id, created_records)

        return JsonResponse({
            'success': True,
            'lote_id': lote_id,
            'total_procesados': len(created_records),
            'message': f"Lote {lote_id} procesado exitosamente con {len(created_records)} solicitudes."
        })
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


def _enviar_correo_lote_masivo(solicitante, tipo, lote_id, records):
    dest = getattr(settings, 'ABASTECIMIENTO_EMAIL', 'abastecimiento@pesco.cl')
    tipo_str = "Creación de Códigos" if tipo == 'codigo' else "Solicitudes de Compra"

    summary_lines = []
    for r in records:
        if tipo == 'codigo':
            summary_lines.append(f" - #{r.pk} [{r.generated_code or 'Pendiente'}]: {r.proposed_description} (Grupo: {r.grupo_material or '-'})")
        else:
            summary_lines.append(f" - #{r.pk} [{r.codigo_compra or 'S/C'}]: {r.descripcion} x {r.cantidad_solicitada} un. (${r.costo_total})")

    resumen_txt = "\n".join(summary_lines[:25])
    if len(records) > 25:
        resumen_txt += f"\n... y {len(records) - 25} ítems más."

    try:
        send_mail(
            subject=f"[PESCO] Nueva Carga Masiva {lote_id} ({len(records)} ítems) - {solicitante}",
            message=(
                f"Se ha ingresado un nuevo LOTE MASIVO de solicitudes en el Portal PESCO.\n\n"
                f"ID de Lote: {lote_id}\n"
                f"Tipo de Lote: {tipo_str}\n"
                f"Solicitante: {solicitante}\n"
                f"Total de Ítems: {len(records)}\n\n"
                f"── RESUMEN DE ÍTEMS EN EL LOTE ──\n"
                f"{resumen_txt}\n\n"
                f"Portal de Administración: http://127.0.0.1:8000/solicitudes/"
            ),
            from_email=settings.DEFAULT_FROM_EMAIL,
            recipient_list=[dest],
            fail_silently=True,
        )
    except Exception:
        pass


